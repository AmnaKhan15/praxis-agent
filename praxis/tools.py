"""The tools the agent can call, and the sandbox that limits them.

The AI model never touches a file itself. It only sends a request such as
"read receipt_07.csv". Our code decides whether to allow it and does the work.
That is the guardrail: the agent can do only what these functions permit.
"""

from __future__ import annotations

import hashlib
import os
import re
import shutil
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import coordinate_from_string

from .config import RECEIPTS_SRC

REPORT_NAME = "expense_report.xlsx"
MAX_READ_CHARS = 6000
BAD_SHEET_CHARS = set('[]:*?/\\')
NUMBER_TEXT = re.compile(r"-?\d{1,3}(,\d{3})*(\.\d+)?|-?\d+(\.\d+)?")


class ToolError(Exception):
    """A tool refused or failed. The message is shown to the agent so it can react."""


# What the model is told about each tool (name, purpose, inputs).
TOOL_SPECS: dict[str, dict] = {
    "list_files": {
        "description": "List the files in a folder of the workspace. Use '.' for the top level or 'receipts' for the receipts.",
        "parameters": {
            "type": "object",
            "properties": {"path": {"type": "string", "description": "Folder to list, e.g. 'receipts'"}},
            "required": ["path"],
        },
    },
    "read_file": {
        "description": "Read a text or csv file, or a .xlsx file (shown as cells). Example path: 'receipts/receipt_01.txt'.",
        "parameters": {
            "type": "object",
            "properties": {"path": {"type": "string", "description": "File to read"}},
            "required": ["path"],
        },
    },
    "write_cell": {
        "description": f"Write one value into one cell of {REPORT_NAME} (created if missing, with a 'Summary' sheet). Write plain values, never formulas.",
        "parameters": {
            "type": "object",
            "properties": {
                "sheet": {"type": "string", "description": "Sheet name, e.g. 'Summary'"},
                "cell": {"type": "string", "description": "Cell, e.g. 'B2'"},
                "value": {"type": "string", "description": "Text or a number, e.g. 'Meals' or '390.85'. Numbers are stored as numbers."},
            },
            "required": ["sheet", "cell", "value"],
        },
    },
    "write_rows": {
        "description": f"Write several rows into {REPORT_NAME} in one call, starting at column A of start_row. Plain values only, no formulas.",
        "parameters": {
            "type": "object",
            "properties": {
                "sheet": {"type": "string", "description": "Sheet name"},
                "start_row": {"type": "integer", "description": "First row number, e.g. 1"},
                "rows": {
                    "type": "array",
                    "description": "List of rows; each row is a list of values. Numbers like 390.85 are stored as numbers.",
                    "items": {"type": "array", "items": {"type": "string"}},
                },
            },
            "required": ["sheet", "start_row", "rows"],
        },
    },
    "add_sheet": {
        "description": f"Add a new empty sheet to {REPORT_NAME}.",
        "parameters": {
            "type": "object",
            "properties": {"name": {"type": "string", "description": "Name of the new sheet, e.g. 'Flags'"}},
            "required": ["name"],
        },
    },
}


def sha256_of(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class Workspace:
    """A private temporary folder holding a fresh copy of the receipts."""

    def __init__(self, root: Path):
        self.root = Path(root).resolve()

    @classmethod
    def create(cls, receipts_src: Path = RECEIPTS_SRC) -> "Workspace":
        """Make a brand-new workspace with an untouched copy of the receipts."""
        root = Path(tempfile.mkdtemp(prefix="praxis_"))
        shutil.copytree(receipts_src, root / "receipts")
        for item in (root / "receipts").iterdir():
            os.chmod(item, 0o444)  # extra safety net; the hash check is the real proof
        return cls(root)

    def cleanup(self) -> None:
        """Delete the workspace (restoring write permission first)."""
        receipts = self.root / "receipts"
        if receipts.exists():
            for item in receipts.iterdir():
                os.chmod(item, 0o644)
        shutil.rmtree(self.root, ignore_errors=True)

    @property
    def report_path(self) -> Path:
        return self.root / REPORT_NAME

    # ---- safety -----------------------------------------------------------
    def _resolve(self, relative: str) -> Path:
        """Turn a path from the agent into a real path, refusing anything outside the workspace."""
        if not isinstance(relative, str) or not relative.strip():
            raise ToolError("A path is required.")
        candidate = Path(relative.strip())
        if candidate.is_absolute():
            raise ToolError("Absolute paths are not allowed. Use a path inside the workspace, e.g. 'receipts/receipt_01.txt'.")
        resolved = (self.root / candidate).resolve()
        if not resolved.is_relative_to(self.root):
            raise ToolError("That path is outside the workspace.")
        return resolved

    # ---- reading ----------------------------------------------------------
    def list_files(self, path: str = ".") -> str:
        folder = self._resolve(path)
        if not folder.is_dir():
            raise ToolError(f"'{path}' is not a folder.")
        entries = sorted(folder.iterdir(), key=lambda p: p.name)
        if not entries:
            return "(empty folder)"
        lines = []
        for entry in entries:
            kind = "folder" if entry.is_dir() else f"{entry.stat().st_size} bytes"
            lines.append(f"{entry.relative_to(self.root).as_posix()}  ({kind})")
        return "\n".join(lines)

    def read_file(self, path: str) -> str:
        target = self._resolve(path)
        if not target.is_file():
            raise ToolError(f"File not found: '{path}'. Use list_files to see what exists.")
        if target.suffix.lower() == ".xlsx":
            return self._dump_workbook(target)
        text = target.read_text(encoding="utf-8", errors="replace")
        if len(text) > MAX_READ_CHARS:
            text = text[:MAX_READ_CHARS] + "\n...(truncated)"
        return text

    @staticmethod
    def _dump_workbook(path: Path) -> str:
        wb = load_workbook(path)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        lines.append(f"  {cell.coordinate} = {cell.value!r}")
        return "\n".join(lines) if lines else "(empty workbook)"

    # ---- writing (only ever the report file) ------------------------------
    def _open_report(self) -> Workbook:
        if self.report_path.exists():
            return load_workbook(self.report_path)
        wb = Workbook()
        wb.active.title = "Summary"
        return wb

    @staticmethod
    def _clean_value(value):
        """Check a value is safe, and store number-looking text as a real number."""
        if isinstance(value, bool) or not isinstance(value, (int, float, str)):
            raise ToolError("Values must be text or numbers.")
        if isinstance(value, str):
            if value.lstrip().startswith("="):
                raise ToolError("Formulas are not supported because they are not calculated. Write the final value instead.")
            if NUMBER_TEXT.fullmatch(value.strip()):
                number = float(value.strip().replace(",", ""))
                return int(number) if number.is_integer() and "." not in value else number
        return value

    @staticmethod
    def _sheet(wb: Workbook, name: str):
        if name not in wb.sheetnames:
            raise ToolError(f"Sheet '{name}' not found. Existing sheets: {', '.join(wb.sheetnames)}.")
        return wb[name]

    def write_cell(self, sheet: str, cell: str, value) -> str:
        value = self._clean_value(value)
        try:
            column, row = coordinate_from_string(str(cell).upper())
        except Exception:
            raise ToolError(f"'{cell}' is not a valid cell like 'B2'.")
        wb = self._open_report()
        self._sheet(wb, sheet)[f"{column}{row}"] = value
        wb.save(self.report_path)
        return f"Wrote {value!r} to {sheet}!{column}{row}"

    def write_rows(self, sheet: str, start_row: int, rows: list) -> str:
        if not isinstance(rows, list) or not rows:
            raise ToolError("rows must be a non-empty list of rows.")
        try:
            start_row = int(start_row)
        except (TypeError, ValueError):
            raise ToolError("start_row must be a whole number.")
        if start_row < 1:
            raise ToolError("start_row must be 1 or more.")
        wb = self._open_report()
        ws = self._sheet(wb, sheet)
        for offset, row in enumerate(rows):
            if not isinstance(row, list):
                raise ToolError("Each row must be a list of values.")
            for col, value in enumerate(row, start=1):
                ws.cell(row=start_row + offset, column=col, value=self._clean_value(value))
        wb.save(self.report_path)
        return f"Wrote {len(rows)} row(s) to {sheet} starting at row {start_row}"

    def add_sheet(self, name: str) -> str:
        if not isinstance(name, str) or not name.strip() or len(name) > 31 or set(name) & BAD_SHEET_CHARS:
            raise ToolError("Sheet names must be 1 to 31 characters and cannot contain []:*?/\\")
        wb = self._open_report()
        if name in wb.sheetnames:
            raise ToolError(f"Sheet '{name}' already exists.")
        wb.create_sheet(name)
        wb.save(self.report_path)
        return f"Added sheet '{name}'. Sheets now: {', '.join(wb.sheetnames)}."

    # ---- dispatcher -------------------------------------------------------
    def call(self, name: str, args: dict, allowed: list[str]) -> str:
        """Run one tool call if the tool is allowed. Raises ToolError with a clear message otherwise."""
        if name not in allowed:
            raise ToolError(f"Tool '{name}' is not available in this workspace. Available tools: {', '.join(allowed)}.")
        if not isinstance(args, dict):
            raise ToolError("Tool arguments must be an object.")
        try:
            return getattr(self, name)(**args)
        except TypeError as exc:
            raise ToolError(f"Wrong arguments for {name}: {exc}")


def receipt_number(text: str) -> int | None:
    """Find a receipt number such as 'receipt_07' inside any text."""
    match = re.search(r"receipt_(\d+)", text.lower())
    return int(match.group(1)) if match else None
