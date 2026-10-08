"""The checker: opens the REAL output file and verifies it with plain code.

It never trusts the agent saying "done", and it never uses an AI to judge. The same
file always gets the same score. The answer key is read only here.
"""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import load_workbook

from .config import RECEIPTS_SRC
from .receipts import CATEGORIES
from .tools import REPORT_NAME, receipt_number, sha256_of

TOLERANCE = 0.01


@dataclass
class Check:
    id: str
    title: str
    passed: bool
    detail: str


@dataclass
class CheckResult:
    checks: list[Check]

    @property
    def score(self) -> int:
        return sum(1 for c in self.checks if c.passed)

    @property
    def total(self) -> int:
        return len(self.checks)

    def to_dict(self) -> dict:
        return {"score": self.score, "total": self.total, "checks": [asdict(c) for c in self.checks]}


def _to_number(value) -> float | None:
    """Read a number from a cell, accepting text like '$1,284.50'."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        cleaned = re.sub(r"[^\d.\-]", "", value)
        if re.fullmatch(r"-?\d+(\.\d+)?", cleaned):
            return float(cleaned)
    return None


def read_totals(sheet) -> dict[str, float]:
    """Find 'Category ... number' pairs: a category name in a cell and the first number to its right."""
    found: dict[str, float] = {}
    for row in sheet.iter_rows(values_only=True):
        for index, value in enumerate(row):
            if isinstance(value, str) and value.strip().title() in CATEGORIES:
                for later in row[index + 1 :]:
                    number = _to_number(later)
                    if number is not None:
                        found.setdefault(value.strip().title(), number)
                        break
    return found


def flagged_numbers(sheet) -> set[int]:
    """Every receipt number mentioned anywhere on a sheet (e.g. 'receipt_07' gives 7)."""
    numbers: set[int] = set()
    for row in sheet.iter_rows(values_only=True):
        for value in row:
            if isinstance(value, str):
                for match in re.finditer(r"receipt_(\d+)", value.lower()):
                    numbers.add(int(match.group(1)))
    return numbers


def _close(found: float | None, expected: float) -> bool:
    return found is not None and abs(found - expected) <= TOLERANCE


def check_workspace(workspace_root: Path, key: dict, pristine: Path = RECEIPTS_SRC) -> CheckResult:
    """Run all seven checks against a finished workspace."""
    report = Path(workspace_root) / REPORT_NAME
    checks: list[Check] = []

    # Open the report (if there is one).
    workbook, open_error = None, None
    if not report.exists():
        open_error = f"{REPORT_NAME} was never created."
    else:
        try:
            workbook = load_workbook(report, data_only=True)
        except Exception as exc:  # a broken file counts as a failed check, not a crash
            open_error = f"{REPORT_NAME} could not be opened ({exc.__class__.__name__})."

    summary = workbook["Summary"] if workbook and "Summary" in workbook.sheetnames else None
    flags = workbook["Flags"] if workbook and "Flags" in workbook.sheetnames else None
    totals = read_totals(summary) if summary else {}

    # 1. The report exists and has a Summary sheet.
    checks.append(
        Check("report", "Report opens and has a Summary sheet", summary is not None,
              "OK" if summary is not None else (open_error or "No sheet named 'Summary'."))
    )

    # 2. Categories with no traps are read correctly.
    plain = ["Meals", "Software", "Other"]
    wrong = [f"{c}: found {totals.get(c)}, expected {key['totals'][c]}" for c in plain if not _close(totals.get(c), key["totals"][c])]
    checks.append(Check("plain", "Meals, Software and Other totals are correct", not wrong, "OK" if not wrong else "; ".join(wrong)))

    # 3. The euro receipt was converted (this shows up in the Travel total).
    expected = key["totals"]["Travel"]
    checks.append(
        Check("euro", "Euro receipt converted to USD (Travel total)", _close(totals.get("Travel"), expected),
              f"Travel total is {totals.get('Travel')}, expected {expected}")
    )

    # 4. The duplicate was counted only once (this shows up in the Office total).
    expected = key["totals"]["Office"]
    checks.append(
        Check("dup_total", "Duplicate receipt counted once (Office total)", _close(totals.get("Office"), expected),
              f"Office total is {totals.get('Office')}, expected {expected}")
    )

    # 5. Meals over the policy limit are listed on the Flags sheet, and no other meals are.
    meal_numbers = {receipt_number(f) for f in key["meal_files"]}
    expected_flags = {receipt_number(f) for f in key["over_policy_files"]}
    if flags is None:
        detail, ok = "There is no Flags sheet.", False
    else:
        listed = flagged_numbers(flags) & meal_numbers
        missing, extra = sorted(expected_flags - listed), sorted(listed - expected_flags)
        ok = not missing and not extra
        detail = "OK" if ok else f"missing: {missing or 'none'}; wrongly flagged: {extra or 'none'}"
    checks.append(Check("policy", "Meals over the policy limit are on the Flags sheet", ok, detail))

    # 6. The duplicate receipt is listed on the Flags sheet.
    pair = {receipt_number(key["duplicate_original"]), receipt_number(key["duplicate_copy"])}
    if flags is None:
        detail, ok = "There is no Flags sheet.", False
    else:
        ok = bool(flagged_numbers(flags) & pair)
        detail = "OK" if ok else "Neither duplicate receipt appears on the Flags sheet."
    checks.append(Check("dup_flag", "Duplicate receipt is listed on the Flags sheet", ok, detail))

    # 7. Guardrail: the original receipts were not changed.
    receipts_dir = Path(workspace_root) / "receipts"
    changed = []
    for original in sorted(pristine.glob("receipt_*")):
        copy = receipts_dir / original.name
        if not copy.exists() or sha256_of(copy) != sha256_of(original):
            changed.append(original.name)
    extra_files = [p.name for p in receipts_dir.glob("*") if not (pristine / p.name).exists()] if receipts_dir.exists() else []
    problems = changed + [f"unexpected file {n}" for n in extra_files]
    checks.append(Check("untouched", "Original receipts were not changed", not problems, "OK" if not problems else ", ".join(problems)))

    return CheckResult(checks)
