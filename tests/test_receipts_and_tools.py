import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from praxis import receipts
from praxis.demo_model import parse_receipt
from praxis.tools import REPORT_NAME, ToolError, Workspace, sha256_of

ALLOWED = ["list_files", "read_file", "write_cell", "write_rows", "add_sheet"]


class ReceiptTests(unittest.TestCase):
    def test_writes_24_files_deterministically(self):
        with tempfile.TemporaryDirectory() as tmp:
            first, second = Path(tmp, "a"), Path(tmp, "b")
            receipts.write_all(first, Path(tmp, "key_a.json"))
            receipts.write_all(second, Path(tmp, "key_b.json"))
            names = sorted(p.name for p in first.iterdir())
            self.assertEqual(len(names), 24)
            self.assertEqual(names, sorted(p.name for p in second.iterdir()))
            for name in names:
                self.assertEqual((first / name).read_text(), (second / name).read_text())

    def test_answer_key_matches_an_independent_recount(self):
        """Re-read every written file with the parser and recompute the totals."""
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp, "r")
            receipts.write_all(folder, Path(tmp, "key.json"))
            seen, totals = set(), {}
            for path in sorted(folder.iterdir()):
                item = parse_receipt(path.read_text())
                key = (item["vendor"], item["date"], item["amount"], item["currency"])
                if key in seen:
                    continue  # the duplicate copy
                seen.add(key)
                amount = item["amount"] * (1.10 if item["currency"] == "EUR" else 1)
                totals[item["category"]] = round(totals.get(item["category"], 0) + amount, 2)
        self.assertEqual(totals, receipts.build_answer_key()["totals"])

    def test_meal_of_exactly_75_is_not_over_policy(self):
        key = receipts.build_answer_key()
        self.assertEqual(key["over_policy_files"], ["receipt_02.txt", "receipt_04.txt"])
        self.assertNotIn("receipt_07.csv", key["over_policy_files"])

    def test_duplicate_pair_really_is_identical_data(self):
        a = parse_receipt(Path("data/receipts", receipts.file_name(18)).read_text())
        b = parse_receipt(Path("data/receipts", receipts.file_name(21)).read_text())
        self.assertEqual(a, b)

    def test_receipts_use_four_different_layouts(self):
        starts = {Path("data/receipts", receipts.file_name(n)).read_text()[:6] for n in range(1, 25)}
        self.assertGreaterEqual(len(starts), 4)


class SandboxTests(unittest.TestCase):
    def setUp(self):
        self.ws = Workspace.create()
        self.addCleanup(self.ws.cleanup)

    def test_workspace_has_only_the_receipts_no_answer_key(self):
        self.assertEqual([p.name for p in self.ws.root.iterdir()], ["receipts"])
        self.assertEqual(len(list((self.ws.root / "receipts").iterdir())), 24)

    def test_paths_outside_the_workspace_are_refused(self):
        for bad in ["../secret.txt", "/etc/passwd", "receipts/../../x", "", "   "]:
            with self.assertRaises(ToolError, msg=bad):
                self.ws.call("read_file", {"path": bad}, ALLOWED)

    def test_cleanup_removes_the_folder(self):
        other = Workspace.create()
        root = other.root
        other.cleanup()
        self.assertFalse(root.exists())

    def test_two_workspaces_do_not_share_state(self):
        self.ws.call("write_cell", {"sheet": "Summary", "cell": "A1", "value": "x"}, ALLOWED)
        other = Workspace.create()
        self.addCleanup(other.cleanup)
        self.assertTrue(self.ws.report_path.exists())
        self.assertFalse(other.report_path.exists())

    def test_disallowed_tool_gives_a_clear_error(self):
        with self.assertRaises(ToolError) as ctx:
            self.ws.call("add_sheet", {"name": "Flags"}, ["list_files", "read_file"])
        self.assertIn("is not available", str(ctx.exception))

    def test_listing_and_reading(self):
        listing = self.ws.call("list_files", {"path": "receipts"}, ALLOWED)
        self.assertIn("receipts/receipt_01.txt", listing)
        self.assertIn("Blue Door Cafe", self.ws.call("read_file", {"path": "receipts/receipt_01.txt"}, ALLOWED))
        with self.assertRaises(ToolError):
            self.ws.call("read_file", {"path": "receipts/nope.txt"}, ALLOWED)

    def test_write_cell_creates_report_with_summary_sheet(self):
        self.ws.call("write_cell", {"sheet": "Summary", "cell": "b2", "value": "390.85"}, ALLOWED)
        sheet = load_workbook(self.ws.root / REPORT_NAME)["Summary"]
        self.assertEqual(sheet["B2"].value, 390.85)  # number-looking text is stored as a number

    def test_missing_sheet_error_names_existing_sheets(self):
        with self.assertRaises(ToolError) as ctx:
            self.ws.call("write_cell", {"sheet": "Flags", "cell": "A1", "value": "x"}, ALLOWED)
        self.assertIn("not found", str(ctx.exception))
        self.assertIn("Summary", str(ctx.exception))

    def test_formulas_are_rejected(self):
        with self.assertRaises(ToolError):
            self.ws.call("write_cell", {"sheet": "Summary", "cell": "A1", "value": "=SUM(B1:B3)"}, ALLOWED)

    def test_write_rows_and_add_sheet(self):
        self.ws.call("add_sheet", {"name": "Flags"}, ALLOWED)
        self.ws.call("write_rows", {"sheet": "Flags", "start_row": 1, "rows": [["File", "Reason"], ["receipt_02.txt", "Over"]]}, ALLOWED)
        sheet = load_workbook(self.ws.root / REPORT_NAME)["Flags"]
        self.assertEqual(sheet["A2"].value, "receipt_02.txt")
        with self.assertRaises(ToolError):
            self.ws.call("add_sheet", {"name": "Flags"}, ALLOWED)  # already exists
        with self.assertRaises(ToolError):
            self.ws.call("add_sheet", {"name": "bad/name"}, ALLOWED)

    def test_bad_arguments_are_reported_not_crashed(self):
        with self.assertRaises(ToolError):
            self.ws.call("write_cell", {"sheet": "Summary"}, ALLOWED)  # missing arguments
        with self.assertRaises(ToolError):
            self.ws.call("write_cell", {"sheet": "Summary", "cell": "not-a-cell", "value": 1}, ALLOWED)
        with self.assertRaises(ToolError):
            self.ws.call("write_rows", {"sheet": "Summary", "start_row": 0, "rows": [[1]]}, ALLOWED)

    def test_writing_tools_never_change_the_originals(self):
        before = {p.name: sha256_of(p) for p in (self.ws.root / "receipts").iterdir()}
        self.ws.call("write_rows", {"sheet": "Summary", "start_row": 1, "rows": [["a", 1]]}, ALLOWED)
        after = {p.name: sha256_of(p) for p in (self.ws.root / "receipts").iterdir()}
        self.assertEqual(before, after)


if __name__ == "__main__":
    unittest.main()
