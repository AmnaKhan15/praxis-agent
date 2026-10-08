import shutil
import tempfile
import unittest
from pathlib import Path

from praxis.checker import check_workspace
from praxis.diagnose import diagnose
from praxis.receipts import load_answer_key
from praxis.tools import REPORT_NAME, Workspace
from helpers import build_report

KEY = load_answer_key()


class CheckerTests(unittest.TestCase):
    def setUp(self):
        self.ws = Workspace.create()
        self.addCleanup(self.ws.cleanup)

    def run_check(self, **kwargs):
        build_report(self.ws.root / REPORT_NAME, KEY, **kwargs)
        return check_workspace(self.ws.root, KEY)

    def failed(self, result):
        return {c.id for c in result.checks if not c.passed}

    def test_a_perfect_report_scores_7_of_7(self):
        result = self.run_check()
        self.assertEqual((result.score, result.total), (7, 7))

    def test_no_report_fails_everything_except_the_guardrail(self):
        result = check_workspace(self.ws.root, KEY)
        self.assertEqual(self.failed(result), {"report", "plain", "euro", "dup_total", "policy", "dup_flag"})

    def test_unconverted_euro_fails_only_the_euro_check(self):
        self.assertEqual(self.failed(self.run_check(travel=530.30)), {"euro"})

    def test_duplicate_counted_twice_fails_only_that_check(self):
        self.assertEqual(self.failed(self.run_check(office=423.90)), {"dup_total"})

    def test_missing_flags_sheet_fails_both_flag_checks(self):
        self.assertEqual(self.failed(self.run_check(flags=None)), {"policy", "dup_flag"})

    def test_wrongly_flagging_a_meal_at_exactly_75_fails(self):
        self.assertEqual(self.failed(self.run_check(extra_meal_flag=True)), {"policy"})

    def test_flagging_the_original_instead_of_the_copy_still_counts(self):
        build_report(self.ws.root / REPORT_NAME, KEY, flags="correct")
        from openpyxl import load_workbook
        wb = load_workbook(self.ws.root / REPORT_NAME)
        wb["Flags"].cell(row=4, column=1, value=KEY["duplicate_original"])
        wb.save(self.ws.root / REPORT_NAME)
        self.assertEqual(check_workspace(self.ws.root, KEY).score, 7)

    def test_totals_written_as_text_are_understood(self):
        from openpyxl import load_workbook
        build_report(self.ws.root / REPORT_NAME, KEY)
        wb = load_workbook(self.ws.root / REPORT_NAME)
        wb["Summary"]["B2"] = "$390.85"
        wb.save(self.ws.root / REPORT_NAME)
        self.assertEqual(check_workspace(self.ws.root, KEY).score, 7)

    def test_a_changed_receipt_fails_the_guardrail(self):
        build_report(self.ws.root / REPORT_NAME, KEY)
        target = self.ws.root / "receipts" / "receipt_01.txt"
        target.chmod(0o644)
        target.write_text("tampered")
        self.assertEqual(self.failed(check_workspace(self.ws.root, KEY)), {"untouched"})

    def test_a_corrupt_report_is_a_failed_check_not_a_crash(self):
        (self.ws.root / REPORT_NAME).write_text("this is not a spreadsheet")
        result = check_workspace(self.ws.root, KEY)
        self.assertIn("report", self.failed(result))


class DiagnoseTests(unittest.TestCase):
    def setUp(self):
        self.ws = Workspace.create()
        self.addCleanup(self.ws.cleanup)

    def levers(self, **kwargs):
        build_report(self.ws.root / REPORT_NAME, KEY, **kwargs)
        return {h.lever for h in diagnose(check_workspace(self.ws.root, KEY), [])}

    def test_a_perfect_run_needs_no_changes(self):
        build_report(self.ws.root / REPORT_NAME, KEY)
        self.assertEqual(diagnose(check_workspace(self.ws.root, KEY), []), [])

    def test_euro_failure_points_at_context(self):
        self.assertEqual(self.levers(travel=530.3), {"context"})

    def test_duplicate_failure_points_at_instructions(self):
        self.assertEqual(self.levers(office=423.9), {"instructions"})

    def test_a_blocked_tool_in_the_trajectory_points_at_workspace(self):
        build_report(self.ws.root / REPORT_NAME, KEY, flags=None)
        steps = [{"n": 7, "type": "tool", "name": "add_sheet", "args": {}, "ok": False,
                  "observation": "ERROR: Tool 'add_sheet' is not available in this workspace."}]
        hints = diagnose(check_workspace(self.ws.root, KEY), steps)
        workspace = [h for h in hints if h.lever == "workspace"]
        self.assertEqual(len(workspace), 1)
        self.assertIn("Step 7", workspace[0].evidence)

    def test_missing_flags_without_a_tool_error_points_at_context(self):
        self.assertIn("context", self.levers(flags=None))


if __name__ == "__main__":
    unittest.main()
