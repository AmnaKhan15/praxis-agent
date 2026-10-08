"""The web API logic, tested without any web framework."""

import base64
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from praxis import caps, web_core as core
from praxis.demo_model import DemoClient
from praxis.llm import LLMError


def collect(req, factory):
    """Run a stream to the end and return [(event, data), ...]."""
    out = []
    for chunk in core.stream_run(req, client_factory=factory):
        if chunk.startswith(":"):
            continue
        event, data = chunk.strip().split("\n", 1)
        out.append((event.removeprefix("event: "), json.loads(data.removeprefix("data: "))))
    return out


class ParseTests(unittest.TestCase):
    def test_defaults_to_demo_and_baseline_flags(self):
        req = core.parse_run_request({})
        self.assertEqual((req["provider"], req["context"], req["workspace"], req["instructions"]), ("demo", False, False, False))

    def test_named_setup_sets_flags(self):
        req = core.parse_run_request({"setup": "instructions"})
        self.assertEqual((req["context"], req["workspace"], req["instructions"]), (True, True, True))

    def test_rejects_bad_input(self):
        bad = [
            "not a dict",
            {"provider": "nope"},
            {"setup": "nope"},
            {"model": "bad model!"},
            {"model": "x" * 200},
            {"api_key": "has space in it 12345"},
            {"api_key": "short"},
            {"overrides": {"goal": "hack"}},
            {"overrides": {"context": "x" * 5000}},
            {"overrides": {"context": 5}},
        ]
        for payload in bad:
            with self.subTest(payload=str(payload)[:40]):
                with self.assertRaises(ValueError):
                    core.parse_run_request(payload)

    def test_accepts_good_key_and_overrides(self):
        req = core.parse_run_request({"provider": "gemini", "api_key": "AIza-abc_123.xyz", "model": "gemini-flash-latest",
                                      "overrides": {"context": "1 EUR = 1.2 USD"}})
        self.assertEqual(req["overrides"], {"context": "1 EUR = 1.2 USD"})


class StreamTests(unittest.TestCase):
    def test_demo_run_streams_steps_then_result_with_report(self):
        req = core.parse_run_request({"setup": "instructions"})
        events = collect(req, lambda _r: DemoClient(delay=0.0))
        kinds = [k for k, _ in events]
        self.assertEqual(kinds[0], "start")
        self.assertEqual(kinds[-1], "result")
        self.assertIn("step", kinds)
        start = events[0][1]
        self.assertTrue(start["demo"])
        result = events[-1][1]
        self.assertEqual(result["record"]["check"]["score"], 7)
        sheet_names = load_workbook(io.BytesIO(base64.b64decode(result["report_b64"]))).sheetnames
        self.assertEqual(sheet_names, ["Summary", "Flags"])

    def test_nothing_is_saved_to_logs(self):
        before = set(Path(core.LOGS_DIR).glob("*"))
        collect(core.parse_run_request({"setup": "baseline"}), lambda _r: DemoClient(delay=0.0))
        self.assertEqual(before, set(Path(core.LOGS_DIR).glob("*")))

    def test_missing_key_is_a_friendly_error_event(self):
        req = core.parse_run_request({"provider": "gemini"})
        with mock.patch.dict(os.environ, {}, clear=False):
            os.environ.pop("GEMINI_API_KEY", None)
            events = collect(req, core.pick_client)
        self.assertEqual(events[0][0], "error")
        self.assertIn("API key", events[0][1]["message"])

    def test_unexpected_crash_does_not_leak_details(self):
        class Boom(DemoClient):
            def chat(self, *a, **k):
                raise RuntimeError("secret-key-SHOULD-NOT-APPEAR")

        events = collect(core.parse_run_request({"setup": "baseline"}), lambda _r: Boom(delay=0.0))
        text = json.dumps(events)
        self.assertNotIn("SHOULD-NOT-APPEAR", text)

    def test_user_key_is_not_echoed_back(self):
        req = core.parse_run_request({"provider": "gemini", "api_key": "AIza-visitor-key-1234"})

        def factory(_r):
            raise LLMError("The provider refused the key.")

        self.assertNotIn("visitor-key", json.dumps(collect(req, factory)))


class CapTests(unittest.TestCase):
    def test_host_key_is_capped_and_counted(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "usage.json"
            req = core.parse_run_request({"provider": "gemini"})
            env = {"GEMINI_API_KEY": "host-key-1234567", "DAILY_RUN_CAP": "1"}
            with mock.patch.dict(os.environ, env):
                core.pick_client(req, path)  # uses the one allowed run
                self.assertEqual(caps.runs_left(path), 0)
                with self.assertRaises(LLMError):
                    core.pick_client(req, path)

    def test_visitor_key_is_not_counted(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "usage.json"
            req = core.parse_run_request({"provider": "gemini", "api_key": "visitor-key-1234567"})
            with mock.patch.dict(os.environ, {"DAILY_RUN_CAP": "1"}):
                core.pick_client(req, path)
                self.assertEqual(caps.runs_today(path), 0)


class DataTests(unittest.TestCase):
    def test_config_never_contains_key_values(self):
        with mock.patch.dict(os.environ, {"GEMINI_API_KEY": "host-key-should-hide"}):
            text = json.dumps(core.public_config())
        self.assertNotIn("should-hide", text)
        self.assertIn('"server_key": true', text)

    def test_examples_and_replay(self):
        listing = core.list_examples()
        self.assertEqual(len(listing), 4)
        self.assertTrue(all(item["demo"] for item in listing))
        found = core.get_example(listing[0]["id"])
        self.assertIn("steps", found["record"])

    def test_example_id_is_strict(self):
        for bad in ["../secret", "example_x/../../etc", "foo", "", "example_demo_baseline.json", "EXAMPLE_x"]:
            self.assertIsNone(core.get_example(bad), bad)

    def test_receipts_listing_has_24(self):
        self.assertEqual(len(core.receipts_listing()), 24)

    def test_answer_key_is_never_served(self):
        blob = json.dumps(core.receipts_listing()) + json.dumps(core.public_config())
        self.assertNotIn("over_policy_files", blob)


if __name__ == "__main__":
    unittest.main()
